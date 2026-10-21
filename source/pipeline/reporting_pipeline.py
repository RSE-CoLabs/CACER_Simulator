from __future__ import annotations

from collections.abc import Callable, Mapping, Sequence
from dataclasses import dataclass
import json
import os
from pathlib import Path
import shutil
from typing import Any

import pandas as pd
import plotly.graph_objects as go
import yaml
from docx import Document
from docx.document import Document as DocumentType

from source.config.utilities import resolve_path
from source.data_structures import (
    ConfigResults,
    EnergyResults,
    FinanceResults,
    GridResults,
    LoadResults,
    PipelineOptions,
    ReportingResults,
    RunContext,
)
from source.reporting.config_visualization import (
    configuration_summary_table,
    registry_summary_table,
    summary_table_figure,
)
from source.reporting.energy_visualization import (
    daily_energy_flows_figure,
    energy_summary_table,
    plot_area_lifetime_consumption,
    plot_area_lifetime_generation,
    plot_area_year1_consumption,
    plot_area_year1_generation,
    plot_pie_consumption,
    plot_pie_generation,
    plot_sankey,
)
from source.reporting.finance_visualization import (
    bill_line_figure,
    bill_pie_figure,
    cash_flow_figure,
    financial_series_figure,
    financial_table_figure,
)
from source.reporting.load_visualization import (
    aggregate_load_table,
    daily_load_profile_figure,
    load_profile_figure,
)


@dataclass(frozen=True, slots=True)
class ReportingData:
    config: ConfigResults
    load: LoadResults
    energy: EnergyResults
    grid: GridResults
    finance: FinanceResults


@dataclass(frozen=True, slots=True)
class ReportingTask:
    name: str
    renderer: Callable[[ReportingData], Any]
    formats: tuple[str, ...] = ("html",)


ACCEPTED_EMPTY_MARKERS = (
    "assets/energy/Energy consumption for a week (weekday, saturday and sunday)",
    "assets/general/1_Yearly energy consumption",
    "assets/general/2_Yearly energy consumption",
    "assets/general/4_Energy consumption for a week (weekday, saturday and sunday)",
    "assets/general/5_Energy consumption for a week (weekday, saturday and sunday)",
    "assets/general/6_Energy consumption for a week (weekday, saturday and sunday)",
)


def run_reporting_pipeline(
    run_context: RunContext,
    config_results: ConfigResults,
    upstream_results: Any | None = None,
    options: PipelineOptions | None = None,
) -> ReportingResults:
    """Render a caller-selected report plan or load explicit external artifacts."""
    selected_options = options or run_context.options.reporting
    if not selected_options.enabled:
        return _disabled_result(selected_options)
    mode = str(selected_options.settings.get("mode", "disabled")).lower()
    if mode == "disabled":
        return _disabled_result(selected_options)
    if mode == "external":
        return _load_external_report(
            run_context,
            selected_options,
        )
    if mode != "execute":
        raise ValueError("Reporting mode must be execute, external, or disabled")

    data = _reporting_data(config_results, upstream_results)
    tasks = selected_options.settings.get("tasks")
    if tasks is None:
        presentation = _load_presentation_inputs(
            run_context,
            selected_options.settings,
        )
        tasks = default_reporting_tasks(data, presentation)
    if not isinstance(tasks, Sequence):
        raise TypeError("Reporting tasks must be a sequence")
    artifacts: list[Path] = []
    figures: dict[str, Any] = {}
    tables: dict[str, Any] = {}
    documents: dict[str, Any] = {}
    for task in tasks:
        if not isinstance(task, ReportingTask):
            raise TypeError("Every reporting task must be a ReportingTask")
        rendered = task.renderer(data)
        saved = save_rendered_object(
            rendered,
            run_context.reporting_output,
            task.name,
            task.formats,
        )
        artifacts.extend(saved)
        if isinstance(rendered, go.Figure):
            figures[task.name] = rendered
        elif isinstance(rendered, pd.DataFrame):
            tables[task.name] = rendered
        elif isinstance(rendered, DocumentType):
            documents[task.name] = rendered
        else:
            raise TypeError(
                f"Unsupported rendered object for task {task.name}: "
                f"{type(rendered).__name__}"
            )
    if selected_options.settings.get("tasks") is None:
        for relative in ACCEPTED_EMPTY_MARKERS:
            marker = run_context.reporting_output / relative
            marker.parent.mkdir(parents=True, exist_ok=True)
            marker.touch()
            artifacts.append(marker)
    artifact_names = tuple(
        path.relative_to(run_context.reporting_output).as_posix()
        for path in artifacts
    )
    return ReportingResults(
        figures=figures,
        tables=tables,
        documents=documents,
        artifacts=tuple(artifacts),
        metadata={
            "domain": "reporting",
            "enabled": True,
            "version": selected_options.version,
            "status": "completed",
            "mode": "execute",
            "task_names": tuple(task.name for task in tasks),
            "artifact_names": artifact_names,
            "plot_count": sum(
                path.suffix.lower() in {".html", ".png"}
                for path in artifacts
            ),
        },
    )


def save_rendered_object(
    rendered: Any,
    destination: Path,
    name: str,
    formats: Sequence[str],
) -> tuple[Path, ...]:
    """Save an object only to the explicit reporting destination."""
    destination.mkdir(parents=True, exist_ok=True)
    artifacts: list[Path] = []
    if isinstance(rendered, go.Figure):
        for output_format in formats:
            path = destination / f"{name}.{output_format}"
            path.parent.mkdir(parents=True, exist_ok=True)
            writable_path = _windows_long_path(path)
            if output_format == "html":
                rendered.write_html(writable_path)
            elif output_format == "png":
                rendered.write_image(writable_path)
            else:
                raise ValueError(f"Unsupported figure format: {output_format}")
            artifacts.append(path)
    elif isinstance(rendered, pd.DataFrame):
        for output_format in formats:
            path = destination / f"{name}.{output_format}"
            path.parent.mkdir(parents=True, exist_ok=True)
            writable_path = _windows_long_path(path)
            if output_format == "csv":
                rendered.to_csv(writable_path)
            elif output_format == "xlsx":
                rendered.to_excel(writable_path)
            else:
                raise ValueError(f"Unsupported table format: {output_format}")
            artifacts.append(path)
    elif isinstance(rendered, DocumentType):
        if tuple(formats) != ("docx",):
            raise ValueError("Documents support only the docx format")
        path = destination / f"{name}.docx"
        path.parent.mkdir(parents=True, exist_ok=True)
        rendered.save(_windows_long_path(path))
        artifacts.append(path)
    return tuple(artifacts)


def _windows_long_path(path: Path) -> Path:
    if os.name != "nt":
        return path
    resolved = str(path.resolve())
    if resolved.startswith("\\\\?\\"):
        return Path(resolved)
    if resolved.startswith("\\\\"):
        return Path(f"\\\\?\\UNC\\{resolved[2:]}")
    return Path(f"\\\\?\\{resolved}")


def build_word_report(
    title: str,
    *,
    tables: Mapping[str, pd.DataFrame] | None = None,
    image_paths: Sequence[Path] = (),
    template_path: Path | None = None,
) -> Document:
    """Build a Word report from explicit tables and image paths."""
    if template_path is not None and not template_path.is_file():
        raise FileNotFoundError(
            f"Report template does not exist: {template_path}"
        )
    document = Document(str(template_path)) if template_path else Document()
    document.add_heading(title, level=0)
    for heading, frame in (tables or {}).items():
        document.add_heading(heading, level=1)
        display = frame.reset_index()
        table = document.add_table(
            rows=1,
            cols=len(display.columns),
        )
        for index, column in enumerate(display.columns):
            table.rows[0].cells[index].text = str(column)
        for row in display.itertuples(index=False):
            cells = table.add_row().cells
            for index, value in enumerate(row):
                cells[index].text = str(value)
    for path in image_paths:
        readable_path = _windows_long_path(path)
        if not readable_path.is_file():
            raise FileNotFoundError(f"Report image does not exist: {path}")
        document.add_picture(str(readable_path))
    return document


def _load_presentation_inputs(
    run_context: RunContext,
    settings: Mapping[str, Any],
) -> dict[str, Any]:
    configured_root = settings.get("presentation_root")
    root = (
        run_context.input_root / "config" / "reporting"
        if configured_root is None
        else resolve_path(run_context.repository_root, configured_root)
    )
    if not root.is_dir():
        raise FileNotFoundError(
            f"Reporting presentation root does not exist: {root}"
        )

    def load_yaml(name: str) -> Mapping[str, Any]:
        path = root / name
        if not path.is_file():
            raise FileNotFoundError(f"Reporting input does not exist: {path}")
        with path.open(encoding="utf-8") as stream:
            value = yaml.safe_load(stream)
        if not isinstance(value, Mapping):
            raise TypeError(f"Reporting input must be a mapping: {path}")
        return value

    template = root / "template" / "Report_template.docx"
    if not template.is_file():
        raise FileNotFoundError(f"Report template does not exist: {template}")
    return {
        "report": load_yaml("report.yml"),
        "layout": load_yaml("layout_charts.yml"),
        "template": template,
        "destination": run_context.reporting_output,
    }


def default_reporting_tasks(
    data: ReportingData,
    presentation: Mapping[str, Any],
) -> tuple[ReportingTask, ...]:
    """Build the accepted native report plan from explicit result objects."""
    load_profiles = data.load.profiles.get("load")
    if not isinstance(load_profiles, pd.DataFrame):
        raise TypeError("Default reporting requires load profile DataFrame")
    aggregate_energy = data.energy.flows.get("aggregate_hourly")
    user_energy = data.energy.flows.get("users")
    if not isinstance(aggregate_energy, pd.DataFrame):
        raise TypeError("Default reporting requires aggregate energy flows")
    if not isinstance(user_energy, Mapping):
        raise TypeError("Default reporting requires per-user energy flows")
    bills = data.finance.bills.get("users")
    cash_flows = data.finance.cash_flows
    if not isinstance(bills, Mapping):
        raise TypeError("Default reporting requires native bill results")
    if not isinstance(cash_flows, Mapping):
        raise TypeError("Default reporting requires native cash-flow results")

    energy_totals = energy_summary_table(aggregate_energy)["total_kwh"]
    days = ("2026-03-24", "2026-06-24", "2026-09-22", "2026-12-18")
    producing_types = {
        name
        for name, values in data.config.registries["user_types"].items()
        if bool(values.get("producing"))
    }
    consuming_user = next(
        (name for name in user_energy if name in producing_types),
        next(iter(user_energy)),
    )
    consuming_frame = user_energy[consuming_user]
    tasks: list[ReportingTask] = []
    colors = presentation["layout"].get("colors", {})

    def add(
        name: str,
        renderer: Callable[[ReportingData], Any],
        formats: tuple[str, ...] = ("html", "png"),
    ) -> None:
        def render(reporting_data: ReportingData) -> Any:
            rendered = renderer(reporting_data)
            if isinstance(rendered, go.Figure):
                rendered.update_layout(
                    template="plotly_white",
                    font={"color": colors.get("blu2", "#1f4e78")},
                )
            return rendered

        tasks.append(ReportingTask(name, render, formats))

    users_table = registry_summary_table(data.config.registries["users"])
    add(
        "assets/general/0_CACER_users_table",
        lambda _data, table=users_table: summary_table_figure(
            table, title="CACER users"
        ),
        ("png",),
    )
    add(
        "assets/general/table_1_CACER 2tabella",
        lambda _data: summary_table_figure(
            configuration_summary_table(data.config.registries["recap"]),
            title="CACER configuration",
        ),
        ("png",),
    )
    add(
        "assets/general/3_Comparison yearly consumption (year 1)",
        lambda _data: load_profile_figure(
            load_profiles.cumsum(), title="Comparison yearly consumption"
        ),
    )
    add(
        "assets/general/7_Monthly energy consumption and seasonal effect per user types",
        lambda _data: load_profile_figure(
            load_profiles.resample("ME").sum(),
            title="Monthly energy consumption and seasonal effect per user types",
        ),
    )
    add(
        "assets/general/8_CACER aggregated monthly energy consumption",
        lambda _data: load_profile_figure(
            load_profiles.sum(axis=1).resample("ME").sum().to_frame("CACER"),
            title="CACER aggregated monthly energy consumption",
        ),
    )
    add(
        "assets/energy/Monthly energy consumption and seasonal effect per user types",
        lambda _data: load_profile_figure(
            load_profiles.resample("ME").sum(),
            title="Monthly energy consumption and seasonal effect per user types",
        ),
        ("png",),
    )
    add(
        "assets/energy/Energy flows - 2026-06-24 - CACER ",
        lambda _data: daily_energy_flows_figure(
            aggregate_energy, "2026-06-24", user="CACER"
        ),
        ("png",),
    )
    add(
        "assets/energy/Energy flows - 2026-12-18 - CACER ",
        lambda _data: daily_energy_flows_figure(
            aggregate_energy, "2026-12-18", user="CACER"
        ),
        ("png",),
    )
    add(
        "assets/energy/CACER Energy consumption allocation - [%]",
        lambda _data: plot_pie_consumption(energy_totals, "CACER"),
        ("png",),
    )
    add(
        "assets/energy/CACER Energy generation allocation - [%]",
        lambda _data: plot_pie_generation(energy_totals, "CACER"),
        ("png",),
    )
    add(
        "assets/gen_pv/9_Plots sun chart 1",
        lambda _data: load_profile_figure(
            data.energy.production["pv"], title="PV production"
        ),
        ("png",),
    )
    add(
        "assets/gen_pv/10_Plots sun chart 2",
        lambda _data: daily_load_profile_figure(
            data.energy.production["pv"],
            "2026-06-24",
            title="PV production - summer day",
        ),
        ("png",),
    )

    energy_renderers: list[tuple[str, Callable[[ReportingData], go.Figure]]] = [
        ("11_Energy flows - CACER", lambda _: plot_sankey(energy_totals, "CACER")),
        (
            "12_CACER Energy consumption allocation - [%]",
            lambda _: plot_pie_consumption(energy_totals, "CACER"),
        ),
        (
            "13_CACER Energy generation allocation - [%]",
            lambda _: plot_pie_generation(energy_totals, "CACER"),
        ),
        (
            "14_CACER Energy consumption allocation - [%]",
            lambda _: plot_pie_consumption(energy_totals, "CACER"),
        ),
        (
            "15_Energy consumption repartition for the first year - CACER",
            lambda _: plot_area_year1_consumption(aggregate_energy, "CACER"),
        ),
        (
            "16_Energy generation repartition - CACER",
            lambda _: plot_area_lifetime_generation(aggregate_energy, "CACER"),
        ),
        (
            "17_Energy generation repartition for the first year - CACER",
            lambda _: plot_area_year1_generation(aggregate_energy, "CACER"),
        ),
    ]
    for day, number in zip(days, range(18, 22)):
        energy_renderers.append(
            (
                f"{number}_Energy flows - {day} - CACER ",
                lambda _, selected=day: daily_energy_flows_figure(
                    aggregate_energy, selected, user="CACER"
                ),
            )
        )
    energy_renderers.extend(
        [
            (
                "22_Energy flows in 4 days for each of the seasons of the year",
                lambda _: load_profile_figure(
                    pd.concat(
                        [
                            aggregate_energy.loc[
                                pd.to_datetime(aggregate_energy.index).normalize()
                                == pd.Timestamp(day)
                            ]
                            for day in days
                        ]
                    ),
                    title="Energy flows in four seasonal days",
                ),
            ),
            ("23_Energy flows - CACER", lambda _: plot_sankey(energy_totals, "CACER")),
            (
                "24_Energy flows prosumer",
                lambda _: plot_sankey(
                    energy_summary_table(consuming_frame)["total_kwh"],
                    consuming_user,
                ),
            ),
        ]
    )
    for day, number in zip(days, range(25, 29)):
        energy_renderers.append(
            (
                f"{number}_Energy flows - {day} - Prosumer {consuming_user} - {consuming_user}",
                lambda _, selected=day: daily_energy_flows_figure(
                    consuming_frame, selected, user=consuming_user
                ),
            )
        )
    energy_renderers.extend(
        [
            (
                f"29_Energy flows in 4 days for each of the seasons of the year user - {consuming_user} - {consuming_user}",
                lambda _: load_profile_figure(
                    pd.concat(
                        [
                            consuming_frame.loc[
                                pd.to_datetime(consuming_frame.index).normalize()
                                == pd.Timestamp(day)
                            ]
                            for day in days
                        ]
                    ),
                    title=f"Energy flows in four seasonal days - {consuming_user}",
                ),
            ),
            (
                f"30_Allocation energy consumption - {consuming_user} - [%]",
                lambda _: plot_pie_consumption(
                    energy_summary_table(consuming_frame)["total_kwh"],
                    consuming_user,
                ),
            ),
            (
                f"31_Allocation energy consumption - {consuming_user} - [%] - {consuming_user}",
                lambda _: plot_pie_generation(
                    energy_summary_table(consuming_frame)["total_kwh"],
                    consuming_user,
                ),
            ),
            (
                "32_Incentives per types and per configurationsCACER",
                lambda _: financial_series_figure(
                    data.finance.incentives["incentives"].monthly,
                    title="Incentives per types and configurations",
                ),
            ),
        ]
    )
    for name, renderer in energy_renderers:
        add(f"assets/energy/{name}", renderer)

    for user, scenarios in bills.items():
        for scenario, frame in scenarios.items():
            if scenario.endswith("_quarterly") or not isinstance(frame, pd.DataFrame):
                continue
            if user == "household_small_1" and scenario != "bau":
                continue
            prefix = f"Bills user {user} in scenario {scenario}"
            add(
                f"assets/bills/{prefix} - line chart",
                lambda _, table=frame, label=user, case=scenario: bill_line_figure(
                    table, user=label, scenario=case
                ),
            )
            add(
                f"assets/bills/{prefix} - pie chart",
                lambda _, table=frame, label=user, case=scenario: bill_pie_figure(
                    table.sum(), user=label, scenario=case
                ),
            )
            if scenario == "pv":
                add(
                    f"assets/bills/{prefix} - bar chart",
                    lambda _, table=frame: cash_flow_figure(
                        table, title="PV bill components"
                    ),
                )

    incentive = data.finance.incentives["incentives"].monthly
    finance_figures = (
        (
            "33_Incentives generation from the different configurations - CACER",
            financial_series_figure(incentive, title="CACER incentives"),
        ),
        (
            "34_Incentives generation from the different configurations",
            financial_series_figure(incentive, title="Configuration incentives"),
        ),
        (
            "35_Aggregated Discounted Cah Flows for user types",
            financial_series_figure(
                cash_flows["project"]["DCF_monthly"],
                columns=("DCF", "DCF_cum"),
                title="Aggregated discounted cash flows",
            ),
        ),
        (
            "36_Cash Flows in the CACER",
            cash_flow_figure(
                cash_flows["project"]["totals"],
                title="Cash flows in the CACER",
            ),
        ),
        (
            "37_Cash Flows in the CACER - Costs",
            cash_flow_figure(
                cash_flows["project"]["totals"][["capex_total", "opex_total"]],
                title="Cash flows in the CACER - Costs",
            ),
        ),
        (
            "38_Cash Flows in the CACER - Revenues",
            cash_flow_figure(
                cash_flows["project"]["totals"][["revenues_total"]],
                title="Cash flows in the CACER - Revenues",
            ),
        ),
        (
            "Input - PZO seasonal (€ per kWh)",
            financial_series_figure(
                data.finance.incentives["rid"].pzo_hourly,
                title="Input - PZO seasonal (€ per kWh)",
            ),
        ),
        (
            "Input - PZO seasonal - year 1 (€ per kWh)",
            financial_series_figure(
                data.finance.incentives["rid"].pzo_hourly,
                title="Input - PZO seasonal - year 1 (€ per kWh)",
            ),
        ),
    )
    for name, figure in finance_figures:
        add(f"assets/finance/{name}", lambda _, value=figure: value)
    for user_id, tables in cash_flows["users"].items():
        if user_id in {"CACER", "social_fund"}:
            continue
        user_type = data.config.registries["users"][user_id]["user_type"]
        title = f"Financial analysis user - {user_type}"
        add(
            f"assets/finance/{title}_{title}",
            lambda _, table=tables["Results"], label=title: financial_table_figure(
                table, title=label
            ),
        )

    report_tables = {
        "CACER - Input data for the simulation": configuration_summary_table(
            data.config.registries["recap"]
        ),
        "General plots": aggregate_load_table(load_profiles),
        "Sun charts": energy_summary_table(
            data.energy.production["pv"]
        ),
        "CACER - Results": energy_summary_table(aggregate_energy),
        "Prosumers - Risultati energetici": energy_summary_table(
            consuming_frame
        ),
        f"Prosumer: {consuming_user}": energy_summary_table(consuming_frame),
        "CACER - Financial results": cash_flows["project"]["Results"],
        "Bills results": bills[consuming_user]["bau"],
    }
    template = Path(presentation["template"])
    destination = Path(presentation["destination"])
    report_images = tuple(
        destination / f"{task.name}.png"
        for task in tasks
        if "png" in task.formats
    )[:36]
    add(
        "assets/report/Report_template",
        lambda _: Document(str(template)),
        ("docx",),
    )
    add(
        "files/results_finance/CER_test",
        lambda _: build_word_report(
            "CER_test financial report",
            tables=report_tables,
            image_paths=report_images,
            template_path=template,
        ),
        ("docx",),
    )
    return tuple(tasks)


def _load_external_report(
    run_context: RunContext,
    options: PipelineOptions,
) -> ReportingResults:
    configured = options.settings.get("external_root")
    if configured is None:
        raise ValueError("external reporting mode requires external_root")
    root = resolve_path(run_context.repository_root, configured)
    if not root.is_dir():
        raise FileNotFoundError(f"External reporting root does not exist: {root}")
    copy_artifacts = bool(options.settings.get("copy_artifacts", True))
    artifacts: list[Path] = []
    manifest_path = options.settings.get("manifest_path")
    if manifest_path is None:
        source_files = [
            path
            for path in sorted(root.rglob("*"))
            if path.is_file() and path.name != ".gitkeep"
        ]
    else:
        manifest_file = resolve_path(
            run_context.repository_root,
            manifest_path,
        )
        manifest = json.loads(manifest_file.read_text(encoding="utf-8"))
        source_files = [
            root / entry["path"]
            for entry in manifest["files"]
            if Path(entry["path"]).name != ".gitkeep"
        ]
        missing = [path for path in source_files if not _io_path(path).is_file()]
        if missing:
            missing_names = ", ".join(str(path) for path in missing)
            raise FileNotFoundError(
                f"Manifest reporting artifacts are missing: {missing_names}"
            )
    if copy_artifacts:
        for source in source_files:
            target = run_context.reporting_output / source.relative_to(root)
            target.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(_io_path(source), _io_path(target))
            artifacts.append(target)
    documents = {}
    for path in source_files:
        if path.suffix.lower() == ".docx":
            documents[path.relative_to(root).as_posix()] = Document(
                str(_io_path(path))
            )
    names = tuple(path.relative_to(root).as_posix() for path in source_files)
    return ReportingResults(
        documents=documents,
        artifacts=tuple(artifacts),
        metadata={
            "domain": "reporting",
            "enabled": True,
            "version": options.version or "legacy",
            "status": "completed",
            "mode": "external",
            "artifact_names": names,
            "plot_count": sum(
                Path(name).suffix.lower() in {".html", ".png"}
                for name in names
            ),
        },
    )


def _reporting_data(
    config: ConfigResults,
    upstream: Any,
) -> ReportingData:
    if not isinstance(upstream, Mapping):
        raise TypeError("Reporting pipeline requires upstream result mapping")
    expected = {
        "load": LoadResults,
        "energy": EnergyResults,
        "grid": GridResults,
        "finance": FinanceResults,
    }
    values = {}
    for name, result_type in expected.items():
        value = upstream.get(name)
        if not isinstance(value, result_type):
            raise TypeError(f"Reporting pipeline requires {result_type.__name__}")
        values[name] = value
    return ReportingData(config=config, **values)


def _disabled_result(options: PipelineOptions) -> ReportingResults:
    return ReportingResults(
        metadata={
            "domain": "reporting",
            "enabled": options.enabled,
            "version": options.version,
            "status": "disabled",
            "mode": "disabled",
        }
    )


def _io_path(path: Path) -> Path:
    if path.is_absolute() and not str(path).startswith("\\\\?\\"):
        return Path(f"\\\\?\\{path}")
    return path
