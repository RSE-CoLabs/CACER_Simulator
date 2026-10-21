from __future__ import annotations

from collections.abc import Callable, Mapping
from dataclasses import dataclass
from pathlib import Path
import shutil
from typing import Any

import pandas as pd
from openpyxl import load_workbook
import yaml

from source.config.utilities import resolve_path
from source.data_structures import (
    BillingResults,
    CashFlowResults,
    ConfigResults,
    EnergyResults,
    FinancialIndicators,
    FinanceResults,
    GridResults,
    IncentiveResults,
    OwnershipResults,
    PipelineOptions,
    RidResults,
    RunContext,
)
from source.finance.billing import calculate_bills
from source.finance.cash_flows import (
    build_finance_template,
    calculate_cash_flows,
)
from source.finance.incentives import calculate_incentives
from source.finance.ownership import calculate_ownership
from source.finance.rid import calculate_rid


FINANCE_STAGE_ORDER = (
    "bills",
    "rid_pzo",
    "aggregations",
    "ownership",
    "incentives",
    "plant_cash_flows",
    "user_cash_flows",
    "project_aggregation",
    "financial_indicators",
    "reporting_results",
)


@dataclass(frozen=True, slots=True)
class FinanceExecutionContext:
    run_context: RunContext
    config: ConfigResults
    energy: EnergyResults
    grid: GridResults | None
    inputs: Mapping[str, Any]


FinanceOperation = Callable[
    [FinanceExecutionContext, Mapping[str, Any]],
    Any,
]


def run_finance_pipeline(
    run_context: RunContext,
    config_results: ConfigResults,
    upstream_results: Any | None = None,
    options: PipelineOptions | None = None,
) -> FinanceResults:
    """Execute finance stages from typed upstream values or explicit external data."""
    selected_options = options or run_context.options.finance
    if not selected_options.enabled:
        return _disabled_result(selected_options)

    settings = selected_options.settings
    mode = str(settings.get("mode", "disabled")).strip().lower()
    if mode == "disabled":
        return _disabled_result(selected_options)
    if mode == "external":
        return load_external_finance_results(
            run_context,
            settings,
            version=selected_options.version,
        )
    if mode != "execute":
        raise ValueError("Finance mode must be execute, external, or disabled")

    energy, grid = _typed_upstream(upstream_results)
    operations = settings.get("operations")
    if operations is not None:
        if not isinstance(operations, Mapping):
            raise TypeError("Finance operations must be a mapping")
        missing = [
            stage
            for stage in FINANCE_STAGE_ORDER
            if not callable(operations.get(stage))
        ]
        if missing:
            raise ValueError(f"Missing finance stage operations: {missing}")
        inputs = settings.get("inputs", {})
        if not isinstance(inputs, Mapping):
            raise TypeError("Finance inputs must be a mapping")
        context = FinanceExecutionContext(
            run_context=run_context,
            config=config_results,
            energy=energy,
            grid=grid,
            inputs=inputs,
        )
        state: dict[str, Any] = {}
        for stage in FINANCE_STAGE_ORDER:
            operation: FinanceOperation = operations[stage]
            state[stage] = operation(context, state)
        return _results_from_state(
            state,
            selected_options.version,
            mode="execute",
        )

    return _run_native_finance(
        run_context,
        config_results,
        energy,
        grid,
        settings,
        version=selected_options.version,
    )


def _run_native_finance(
    run_context: RunContext,
    config: ConfigResults,
    energy: EnergyResults,
    grid: GridResults | None,
    settings: Mapping[str, Any],
    *,
    version: str | None,
) -> FinanceResults:
    inputs = _load_native_inputs(run_context, settings)
    user_flows = energy.flows.get("users")
    if not isinstance(user_flows, Mapping):
        raise TypeError("EnergyResults.flows['users'] must be a mapping")

    bills = calculate_bills(
        energy_user_flows=user_flows,
        user_types=config.registries["user_types"],
        recap=config.registries["recap"],
        calendar_frame=config.calendars["calendar"],
        market_assumptions=inputs["market"],
        pun_monthly=inputs["pun"],
        market_scenario=config.configuration["market_scenario"],
        delta_t=config.configuration["delta_t"],
    )
    rid = calculate_rid(
        energy_user_flows=user_flows,
        user_types=config.registries["user_types"],
        recap=config.registries["recap"],
        municipalities=config.data_structures["municipalities"],
        province=config.configuration["provincia_it"],
        start_date=config.configuration["start_date"],
        project_lifetime_years=config.configuration["project_lifetime_yrs"],
        pzo_input=inputs["pzo"],
        rid_assumptions=inputs["rid"],
        bt_losses=config.configuration["perdite_BT"],
        mt_losses=config.configuration["perdite_MT"],
    )
    ownership = calculate_ownership(
        users=config.registries["users"],
        plants=config.registries["plants"],
        membership_matrix=config.data_structures["membership_matrix"],
        plant_operation_matrix=config.data_structures[
            "plant_operation_matrix"
        ],
        monthly_calendar=config.calendars["monthly_calendar"],
        funding_schemes=inputs["funding"],
        scenario=inputs["scenario"],
        repartition_schemes=inputs["repartition"],
        surplus_schemes=inputs["surplus"],
        capex_table=inputs["capex"],
        scale_factors=inputs["scale_factors"],
    )
    incentives = calculate_incentives(
        pzo_hourly=rid.pzo_hourly,
        shared_energy_hourly=energy.incentives["hourly"],
        valorization_hourly=energy.incentives["valorization_hourly"],
        shared_energy_yearly=energy.incentives["yearly"],
        plants=config.registries["plants"],
        users=config.registries["users"],
        recap=config.registries["recap"],
        province=config.configuration["provincia_it"],
        municipalities=config.data_structures["municipalities"],
        monthly_calendar=config.calendars["monthly_calendar"],
        incentive_repartition=ownership.repartition["incentives"],
        surplus_repartition=ownership.repartition["surplus"],
        cacer_fee_assumptions=inputs["cacer_fees"],
        trase=config.configuration["TRASe"],
        btau=config.configuration["BTAU"],
        cpr_bt=config.configuration["Cpr_bt"],
        cpr_mt=config.configuration["Cpr_mt"],
    )
    template = build_finance_template(
        config.calendars["monthly_calendar"],
        inputs["inflation"],
        inputs["discount_rates"],
    )
    cash_flows, indicators = calculate_cash_flows(
        template=template,
        users=config.registries["users"],
        plants=config.registries["plants"],
        recap=config.registries["recap"],
        ownership=ownership,
        bills=bills,
        rid=rid,
        incentives=incentives,
        user_entry_matrix=config.data_structures["user_entry_matrix"].drop(
            index="month", errors="ignore"
        ),
        membership_matrix=config.data_structures["membership_matrix"].drop(
            index="month", errors="ignore"
        ),
        plant_operation_matrix=config.data_structures[
            "plant_operation_matrix"
        ],
        capex_table=inputs["capex"],
        opex_plant=inputs["opex_plant"],
        opex_cacer=inputs["opex_cacer"],
        taxes=inputs["taxes"],
        active_cacer_kickoff=inputs["active_cacer_kickoff"],
        entry_fee=inputs["entry_fee"],
        subscription_fee=inputs["subscription_fee"],
        ground_opex_factor=inputs["ground_opex_factor"],
    )
    artifacts = _persist_native_results(
        run_context.finance_output,
        bills,
        rid,
        ownership,
        incentives,
        cash_flows,
    )
    return FinanceResults(
        bills={
            "users": bills.user_bills,
            "aggregated": bills.aggregated_bills,
        },
        cash_flows={
            "plants": cash_flows.plants,
            "users": cash_flows.users,
            "configurations": cash_flows.configurations,
            "project": cash_flows.project,
        },
        incentives={
            "rid": rid,
            "incentives": incentives,
        },
        indicators={
            "ownership": ownership,
            "financial": indicators,
            "grid": grid,
        },
        artifacts=artifacts,
        metadata={
            "domain": "finance",
            "enabled": True,
            "version": version or "native",
            "status": "completed",
            "mode": "execute",
            "stage_order": FINANCE_STAGE_ORDER,
        },
    )


def _load_native_inputs(
    run_context: RunContext,
    settings: Mapping[str, Any],
) -> dict[str, Any]:
    configured_root = settings.get("input_root")
    root = (
        run_context.input_root / "finance"
        if configured_root is None
        else resolve_path(run_context.repository_root, configured_root)
    )
    if not root.is_dir():
        raise FileNotFoundError(f"Finance input root does not exist: {root}")

    workbook = root / "inputs_FM.xlsx"
    with (root / "mercato.yml").open(encoding="utf-8") as stream:
        market = yaml.safe_load(stream)
    with (root / "RID_input.yml").open(encoding="utf-8") as stream:
        rid = yaml.safe_load(stream)
    with (root / "CER_fees.yml").open(encoding="utf-8") as stream:
        cacer_fees = yaml.safe_load(stream)

    funding = pd.read_excel(
        workbook, sheet_name="Funding scheme", index_col=0
    )
    scenario = pd.read_excel(
        workbook, sheet_name="Scenario", index_col=0
    )["Value"].to_dict()
    repartition = pd.read_excel(
        workbook, sheet_name="Repartition", index_col=0
    )
    surplus = pd.read_excel(workbook, sheet_name="Surplus", index_col=0)
    capex = pd.read_excel(
        workbook, sheet_name="CAPEX", usecols="A:G", nrows=6, index_col=0
    ).drop(columns="Unit")
    scale_factors = pd.read_excel(
        workbook, sheet_name="CAPEX", usecols="A:G", skiprows=8, nrows=9
    )
    structure = pd.read_excel(
        workbook, sheet_name="CAPEX", usecols="J:Q", nrows=3, index_col=0
    ).drop(columns="Unit.1")
    opex = pd.read_excel(workbook, sheet_name="OPEX")
    workbook_values = load_workbook(workbook, read_only=True, data_only=True)
    active_cacer_kickoff = _value_right_of_label(
        workbook_values["CAPEX"], "active_cacer_kickoff"
    )
    entry_fee = _value_right_of_label(
        workbook_values["CAPEX"], "entry_fee"
    )
    subscription_fee = _value_right_of_label(
        workbook_values["OPEX"], "subscription_fee"
    )
    ground_opex_factor = _value_right_of_label(
        workbook_values["OPEX"], "Scale factor for ground-mounted PV"
    )
    workbook_values.close()
    return {
        "market": market,
        "rid": rid,
        "cacer_fees": cacer_fees,
        "pun": pd.read_csv(root / "PUN_input_data.csv", index_col="month_number"),
        "pzo": pd.read_csv(root / "hourly_PZO_input.csv"),
        "funding": funding,
        "scenario": scenario,
        "repartition": repartition,
        "surplus": surplus,
        "capex": capex,
        "scale_factors": scale_factors,
        "inflation": pd.read_excel(
            workbook, sheet_name="Inflation", index_col=0
        )["inflation_rate_pa"],
        "discount_rates": structure.loc["discount_rate_pa"].astype(
            float
        ).to_dict(),
        "opex_plant": opex.set_index("opex_plant_table")[
            "Value"
        ].dropna(),
        "opex_cacer": opex.set_index("opex_CACER_table")[
            "Value.1"
        ].dropna(),
        "taxes": pd.read_excel(
            workbook, sheet_name="Taxes", index_col="Item"
        ),
        "active_cacer_kickoff": active_cacer_kickoff,
        "entry_fee": entry_fee,
        "subscription_fee": subscription_fee,
        "ground_opex_factor": ground_opex_factor,
    }


def _value_right_of_label(sheet: Any, label: str) -> float:
    normalized = label.strip().lower()
    for row in sheet.iter_rows():
        for position, cell in enumerate(row):
            if str(cell.value).strip().lower() != normalized:
                continue
            for candidate in row[position + 1 :]:
                if isinstance(candidate.value, (int, float)):
                    return float(candidate.value)
    raise ValueError(f"Missing numeric finance assumption: {label}")


def _persist_native_results(
    root: Path,
    bills: BillingResults,
    rid: RidResults,
    ownership: OwnershipResults,
    incentives: IncentiveResults,
    cash_flows: CashFlowResults,
) -> tuple[Path, ...]:
    artifacts: list[Path] = []

    def write_csv(frame: pd.DataFrame, relative: str) -> None:
        path = root / relative
        path.parent.mkdir(parents=True, exist_ok=True)
        frame.to_csv(path)
        artifacts.append(path)

    def write_workbook(
        tables: Mapping[str, pd.DataFrame],
        relative: str,
        *,
        transpose: bool = True,
    ) -> None:
        path = root / relative
        path.parent.mkdir(parents=True, exist_ok=True)
        with pd.ExcelWriter(path, engine="openpyxl") as writer:
            for sheet, frame in tables.items():
                output = frame.T if transpose else frame
                output.to_excel(writer, sheet_name=sheet)
        artifacts.append(path)

    for user_type, tables in bills.user_bills.items():
        write_workbook(
            tables,
            f"finance/bills/{user_type}.xlsx",
            transpose=False,
        )
    for group, tables in bills.aggregated_bills.items():
        write_workbook(
            tables,
            f"finance/bills/{group}.xlsx",
            transpose=False,
        )

    write_csv(rid.pzo_hourly, "PZO/output_PZO_data.csv")
    write_csv(rid.monthly_revenue, "RID/output_RID.csv")
    write_csv(rid.gse_fees, "RID/output_GSE_RID_fees.csv")
    write_csv(rid.aggregated_revenue, "RID/output_RID_active_CACER.csv")
    write_csv(
        incentives.monthly,
        "incentives/CACER_incentivi.csv",
    )
    write_csv(
        incentives.by_configuration,
        "incentives/CACER_incentivi_per_configuration.csv",
    )
    write_csv(
        ownership.subscription,
        "finance/subscription_matrix.csv",
    )
    write_csv(ownership.investment, "finance/investment_matrix.csv")
    ownership_path = root / "finance/ownership_matrix.xlsx"
    ownership_path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(ownership_path, engine="openpyxl") as writer:
        for plant, frame in ownership.ownership.items():
            frame.to_excel(writer, sheet_name=plant)
    artifacts.append(ownership_path)
    repartition_path = root / "finance/repartition_matrix.xlsx"
    with pd.ExcelWriter(repartition_path, engine="openpyxl") as writer:
        for item, frame in ownership.repartition.items():
            frame.to_excel(writer, sheet_name=item)
    artifacts.append(repartition_path)

    for plant, tables in cash_flows.plants.items():
        write_workbook(tables, f"finance/plants/{plant}.xlsx")
    for user, tables in cash_flows.users.items():
        write_workbook(tables, f"finance/users/{user}.xlsx")
    for configuration, tables in cash_flows.configurations.items():
        write_workbook(
            tables,
            f"finance/configurations/{configuration}.xlsx",
        )
    write_workbook(
        cash_flows.project,
        "finance/configurations/project.xlsx",
    )
    return tuple(artifacts)


def load_external_finance_results(
    run_context: RunContext,
    settings: Mapping[str, Any],
    *,
    version: str | None,
) -> FinanceResults:
    """Load explicitly requested legacy finance artifacts for migration or bypass."""
    configured_root = settings.get("external_root")
    if configured_root is None:
        raise ValueError("external mode requires external_root")
    root = resolve_path(run_context.repository_root, configured_root)
    if not root.is_dir():
        raise FileNotFoundError(f"External finance root does not exist: {root}")

    tables = _load_external_tables(root)
    workbooks = _load_external_workbook_schemas(root)
    workbook_tables = _load_external_workbook_tables(root)
    artifacts = []
    if bool(settings.get("copy_artifacts", True)):
        destination = run_context.finance_output / "external"
        destination.mkdir(exist_ok=True)
        for path in sorted(root.rglob("*")):
            if not path.is_file() or path.name == ".gitkeep":
                continue
            target = destination / path.relative_to(root)
            target.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(path, target)
            artifacts.append(target)

    return FinanceResults(
        bills={
            key: value
            for key, value in workbooks.items()
            if key.startswith("bills/")
        },
        cash_flows={
            key: value
            for key, value in workbooks.items()
            if key.startswith(("users/", "plants/", "configurations/"))
        },
        incentives={
            key: value
            for key, value in tables.items()
            if key.startswith("incentives/")
        },
        indicators={
            "tables": tables,
            "workbook_schemas": workbooks,
            "workbook_tables": workbook_tables,
        },
        artifacts=tuple(artifacts),
        metadata={
            "domain": "finance",
            "enabled": True,
            "version": version or "legacy",
            "status": "completed",
            "mode": "external",
            "external_root": str(root),
            "stage_order": FINANCE_STAGE_ORDER,
        },
    )


def _typed_upstream(
    upstream_results: Any,
) -> tuple[EnergyResults, GridResults | None]:
    if not isinstance(upstream_results, Mapping):
        raise TypeError(
            "Finance pipeline requires upstream energy/grid result mapping"
        )
    energy = upstream_results.get("energy")
    grid = upstream_results.get("grid")
    if not isinstance(energy, EnergyResults):
        raise TypeError("Finance pipeline requires EnergyResults")
    if grid is not None and not isinstance(grid, GridResults):
        raise TypeError("Finance grid input must be GridResults or None")
    return energy, grid


def _results_from_state(
    state: Mapping[str, Any],
    version: str | None,
    *,
    mode: str,
) -> FinanceResults:
    reporting = state["reporting_results"]
    artifacts = ()
    if isinstance(reporting, Mapping):
        artifacts = tuple(reporting.get("artifacts", ()))
    return FinanceResults(
        bills={"bills": state["bills"]},
        cash_flows={
            "plants": state["plant_cash_flows"],
            "users": state["user_cash_flows"],
            "project": state["project_aggregation"],
            "reporting": reporting,
        },
        incentives={
            "rid_pzo": state["rid_pzo"],
            "aggregations": state["aggregations"],
            "incentives": state["incentives"],
        },
        indicators={
            "ownership": state["ownership"],
            "financial": state["financial_indicators"],
        },
        artifacts=artifacts,
        metadata={
            "domain": "finance",
            "enabled": True,
            "version": version or "legacy",
            "status": "completed",
            "mode": mode,
            "stage_order": FINANCE_STAGE_ORDER,
        },
    )


def _disabled_result(options: PipelineOptions) -> FinanceResults:
    return FinanceResults(
        metadata={
            "domain": "finance",
            "enabled": options.enabled,
            "version": options.version,
            "status": "disabled",
            "mode": "disabled",
        }
    )


def _load_external_tables(root: Path) -> dict[str, pd.DataFrame]:
    tables: dict[str, pd.DataFrame] = {}
    for path in sorted(root.rglob("*.csv")):
        tables[path.relative_to(root).as_posix()] = pd.read_csv(path)
    return tables


def _load_external_workbook_schemas(
    root: Path,
) -> dict[str, dict[str, tuple[int, int]]]:
    schemas: dict[str, dict[str, tuple[int, int]]] = {}
    for path in sorted(root.rglob("*.xlsx")):
        if path.stat().st_size == 0:
            continue
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheets = {
            sheet.title: (sheet.max_row, sheet.max_column)
            for sheet in workbook.worksheets
        }
        workbook.close()
        schemas[path.relative_to(root).as_posix()] = sheets
    return schemas


def _load_external_workbook_tables(
    root: Path,
) -> dict[str, pd.DataFrame]:
    selected_sheets = {
        "Capex",
        "Totals",
        "totals",
        "DCF_monthly",
        "DCF_yearly",
        "Results",
    }
    tables: dict[str, pd.DataFrame] = {}
    for path in sorted(root.rglob("*.xlsx")):
        if path.stat().st_size == 0:
            continue
        workbook = pd.ExcelFile(path)
        relative = path.relative_to(root).as_posix()
        for sheet in workbook.sheet_names:
            if sheet not in selected_sheets:
                continue
            tables[f"{relative}::{sheet}"] = pd.read_excel(
                workbook,
                sheet_name=sheet,
                index_col=0,
            )
    return tables
